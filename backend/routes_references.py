from fastapi import APIRouter, HTTPException, UploadFile, File, Request
from typing import Dict, Any
import pandas as pd
from io import BytesIO
from datetime import datetime
import uuid

router = APIRouter(prefix="/api")
db = None


def set_db(database):
    global db
    db = database


@router.post("/references/import-file")
async def import_references_from_file(request: Request, file: UploadFile = File(...)):
    """استيراد مراجع من PDF أو Excel"""
    try:
        from core import authz as _authz
        actor = await _authz.resolve_request_actor(request)
        if actor.role != "admin":
            raise HTTPException(status_code=403, detail={"error": "admin_only_import"})
        contents = await file.read()
        imported_counts = {}

        # Determine file type
        if file.filename.lower().endswith((".xlsx", ".xls")):
            # Import from Excel
            excel_data = pd.read_excel(BytesIO(contents), sheet_name=None)

            # Import DTC Codes
            if "DTC Codes" in excel_data:
                dtc_df = excel_data["DTC Codes"]
                dtc_count = 0

                for _, row in dtc_df.iterrows():
                    dtc_doc = {
                        "id": str(uuid.uuid4()),
                        "type": "dtc",
                        "code": str(row.get("الكود", "")).strip().upper(),
                        "nameAr": str(row.get("الاسم بالعربية", "")),
                        "nameEn": str(row.get("الاسم English", "")),
                        "vehicle": str(row.get("السيارة", "")),
                        "causes": (
                            str(row.get("الأسباب", "")).split("\n")
                            if pd.notna(row.get("الأسباب"))
                            else []
                        ),
                        "fixes": (
                            str(row.get("طرق الإصلاح", "")).split("\n")
                            if pd.notna(row.get("طرق الإصلاح"))
                            else []
                        ),
                        "notes": str(row.get("الملاحظات", "")),
                        "pageNumber": str(row.get("الصفحة", "")),
                        "relatedCodes": (
                            str(row.get("أكواد مشابهة", "")).split(",")
                            if pd.notna(row.get("أكواد مشابهة"))
                            else []
                        ),
                        "source": file.filename,
                        "createdAt": datetime.utcnow(),
                    }

                    # Check if already exists
                    existing = await db.dtc_references.find_one(
                        {"code": dtc_doc["code"], "vehicle": dtc_doc["vehicle"]}
                    )
                    if existing:
                        await db.dtc_references.update_one(
                            {"_id": existing["_id"]}, {"$set": dtc_doc}
                        )
                    else:
                        await db.dtc_references.insert_one(dtc_doc)
                    dtc_count += 1

                imported_counts["dtc"] = dtc_count

            # Import Electrical Components
            if "Electrical Components" in excel_data:
                elec_df = excel_data["Electrical Components"]
                elec_count = 0

                for _, row in elec_df.iterrows():
                    elec_doc = {
                        "id": str(uuid.uuid4()),
                        "type": "electrical",
                        "componentAr": str(row.get("المكون", "")),
                        "componentEn": str(row.get("Component", "")),
                        "voltageNormal": float(row.get("الجهد الطبيعي", 0)),
                        "voltageMin": float(row.get("الجهد الأدنى", 0)),
                        "voltageMax": float(row.get("الجهد الأعلى", 0)),
                        "unit": str(row.get("الوحدة", "V")),
                        "measurementMethod": str(row.get("طريقة القياس", "")),
                        "notes": str(row.get("الملاحظات", "")),
                        "source": file.filename,
                        "createdAt": datetime.utcnow(),
                    }

                    await db.electrical_references.insert_one(elec_doc)
                    elec_count += 1

                imported_counts["electrical"] = elec_count

            # Import Vehicle Specs
            if "Vehicle Specs" in excel_data:
                veh_df = excel_data["Vehicle Specs"]
                veh_count = 0

                for _, row in veh_df.iterrows():
                    veh_doc = {
                        "id": str(uuid.uuid4()),
                        "type": "vehicle_spec",
                        "vehicle": str(row.get("السيارة", "")),
                        "engine": str(row.get("المحرك", "")),
                        "year": str(row.get("السنة", "")),
                        "displacement": str(row.get("السعة", "")),
                        "power": str(row.get("القوة", "")),
                        "torque": str(row.get("العزم", "")),
                        "fuelSystem": str(row.get("نظام الوقود", "")),
                        "ignitionSystem": str(row.get("نظام الإشعال", "")),
                        "notes": str(row.get("الملاحظات", "")),
                        "source": file.filename,
                        "createdAt": datetime.utcnow(),
                    }

                    await db.vehicle_references.insert_one(veh_doc)
                    veh_count += 1

                imported_counts["vehicles"] = veh_count

        elif file.filename.lower().endswith(".pdf"):
            # Import from PDF - fast batch processing
            from PyPDF2 import PdfReader
            import re
            from pathlib import Path

            # Save temporarily
            temp_path = Path("/tmp") / file.filename
            with open(temp_path, "wb") as f:
                f.write(contents)

            # Extract text - limit to first 20 pages for speed
            reader = PdfReader(str(temp_path))
            dtc_count = 0
            max_pages = min(20, len(reader.pages))

            # Batch collect all codes first
            all_codes_with_context = {}
            dtc_pattern = re.compile(
                r"\b(P[0-9A-F]{4}|U[0-9A-F]{4}|C[0-9A-F]{4}|B[0-9A-F]{4})\b",
                re.IGNORECASE,
            )

            for page_num in range(max_pages):
                text = reader.pages[page_num].extract_text()
                codes = set(dtc_pattern.findall(text.upper()))

                for code in codes:
                    if code not in all_codes_with_context:
                        code_idx = text.upper().find(code)
                        if code_idx != -1:
                            start = max(0, code_idx - 100)
                            end = min(len(text), code_idx + 200)
                            all_codes_with_context[code] = {
                                "context": text[start:end],
                                "page": page_num + 1,
                            }

            # Batch insert (faster)
            batch_inserts = []
            for code, info in list(all_codes_with_context.items())[:30]:  # Max 30 codes
                # Check if exists
                existing = await db.dtc_references.find_one(
                    {"code": code, "source": file.filename}
                )
                if existing:
                    continue

                # Extract name quickly
                lines = info["context"].split("\n")
                name = code
                for i, line in enumerate(lines):
                    if code in line.upper() and i + 1 < len(lines):
                        name = lines[i + 1].strip()[:80]
                        break

                batch_inserts.append(
                    {
                        "id": str(uuid.uuid4()),
                        "type": "dtc",
                        "code": code,
                        "nameAr": name,
                        "nameEn": name,
                        "vehicle": extract_vehicle_from_filename(file.filename),
                        "causes": [],
                        "fixes": [],
                        "notes": info["context"][:200],
                        "pageNumber": str(info["page"]),
                        "relatedCodes": [],
                        "source": file.filename,
                        "createdAt": datetime.utcnow(),
                    }
                )

            # Batch insert for speed
            if batch_inserts:
                await db.dtc_references.insert_many(batch_inserts)
                dtc_count = len(batch_inserts)

            imported_counts["dtc_from_pdf"] = dtc_count

        return {
            "status": "ok",
            "message": "تم استيراد المراجع بنجاح",
            "imported": imported_counts,
            "filename": file.filename,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def extract_vehicle_from_filename(filename: str) -> str:
    """استخراج اسم السيارة من اسم الملف"""
    filename_lower = filename.lower()
    if "land cruiser" in filename_lower or "landcruiser" in filename_lower:
        if "200" in filename:
            return "Toyota Land Cruiser 200"
        return "Toyota Land Cruiser"
    elif "hilux" in filename_lower:
        if "1kd" in filename_lower or "2kd" in filename_lower:
            return "Toyota Hilux 1KD/2KD"
        return "Toyota Hilux"
    elif "innova" in filename_lower:
        return "Toyota Innova"
    return "عام"


@router.get("/references/dtc")
async def get_dtc_references(code: str = None, vehicle: str = None):
    """الحصول على مراجع DTC"""
    try:
        query = {}
        if code:
            query["code"] = code.upper()
        if vehicle:
            query["vehicle"] = {"$regex": vehicle, "$options": "i"}

        refs = await db.dtc_references.find(query).to_list(length=100)
        for ref in refs:
            ref.pop("_id", None)

        return {"references": refs, "count": len(refs)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/references/download-excel-program")
async def download_excel_search_program():
    """تحميل برنامج Excel للبحث في المراجع"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from fastapi.responses import StreamingResponse
        import io

        wb = Workbook()

        # Sheet 1: البحث الذكي
        ws_search = wb.active
        ws_search.title = "البحث الذكي"

        # Instructions
        ws_search["A1"] = "برنامج البحث في المراجع الفنية"
        ws_search["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        ws_search["A1"].fill = PatternFill(
            start_color="1bdbac", end_color="1bdbac", fill_type="solid"
        )
        ws_search.merge_cells("A1:F1")

        ws_search["A3"] = "ابحث هنا:"
        ws_search["B3"] = ""  # Search box
        ws_search["B3"].fill = PatternFill(
            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
        )

        ws_search["A5"] = "التعليمات:"
        ws_search["A6"] = "1. اكتب اسم المكون أو الكود في الخلية B3"
        ws_search["A7"] = "2. انظر للنتائج في الأوراق الأخرى"
        ws_search["A8"] = "3. استخدم Ctrl+F للبحث السريع"

        # Sheet 2: جميع أكواد DTC
        ws_dtc = wb.create_sheet("أكواد الأعطال")

        # Get all DTC references
        dtc_refs = await db.dtc_references.find({}).to_list(length=1000)

        headers = [
            "الكود",
            "الاسم",
            "السيارة",
            "الأسباب",
            "الحلول",
            "الصفحة",
            "أكواد مشابهة",
        ]
        ws_dtc.append(headers)

        for cell in ws_dtc[1]:
            cell.fill = PatternFill(
                start_color="111827", end_color="111827", fill_type="solid"
            )
            cell.font = Font(bold=True, color="FFFFFF")

        for dtc in dtc_refs:
            causes_text = (
                "\n".join(dtc.get("causes", []))
                if isinstance(dtc.get("causes"), list)
                else dtc.get("causes", "")
            )
            fixes_text = (
                "\n".join(dtc.get("fixes", []))
                if isinstance(dtc.get("fixes"), list)
                else dtc.get("fixes", "")
            )
            related_text = (
                ", ".join(dtc.get("relatedCodes", []))
                if isinstance(dtc.get("relatedCodes"), list)
                else ""
            )

            ws_dtc.append(
                [
                    dtc.get("code"),
                    dtc.get("nameAr"),
                    dtc.get("vehicle"),
                    causes_text,
                    fixes_text,
                    dtc.get("pageNumber"),
                    related_text,
                ]
            )

        # Sheet 3: المكونات الكهربائية
        ws_elec = wb.create_sheet("الجهد الكهربائي")

        elec_refs = await db.electrical_references.find({}).to_list(length=1000)

        headers_elec = [
            "المكون",
            "Component",
            "الجهد الطبيعي",
            "الأدنى",
            "الأعلى",
            "الوحدة",
            "طريقة القياس",
            "الحالة",
        ]
        ws_elec.append(headers_elec)

        for cell in ws_elec[1]:
            cell.fill = PatternFill(
                start_color="1bdbac", end_color="1bdbac", fill_type="solid"
            )
            cell.font = Font(bold=True, color="FFFFFF")

        for elec in elec_refs:
            ws_elec.append(
                [
                    elec.get("componentAr"),
                    elec.get("componentEn"),
                    elec.get("voltageNormal"),
                    elec.get("voltageMin"),
                    elec.get("voltageMax"),
                    elec.get("unit"),
                    elec.get("measurementMethod"),
                    elec.get("notes"),
                ]
            )

        # Adjust columns
        for ws in [ws_dtc, ws_elec]:
            for col in ws.columns:
                max_length = 0
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[col[0].column_letter].width = min(
                    max_length + 2, 50
                )

        # Save to bytes
        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        return StreamingResponse(
            excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=Technical_Search_Program.xlsx"
            },
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/references/electrical")
async def get_electrical_references(component: str = None):
    """الحصول على مراجع كهربائية"""
    try:
        query = {}
        if component:
            query["$or"] = [
                {"componentAr": {"$regex": component, "$options": "i"}},
                {"componentEn": {"$regex": component, "$options": "i"}},
            ]

        refs = await db.electrical_references.find(query).to_list(length=100)
        for ref in refs:
            ref.pop("_id", None)

        return {"references": refs, "count": len(refs)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/references/electrical/smart-search")
async def smart_search_electrical(payload: Dict[str, Any]):
    """بحث ذكي سريع عن الجهد الكهربائي"""
    try:
        import os
        from emergentintegrations.llm.chat import LlmChat, UserMessage
        import uuid

        query = payload.get("query", "")

        if not query:
            raise HTTPException(status_code=422, detail="query required")

        # Quick keyword matching first (no AI)
        all_components = await db.electrical_references.find({}).to_list(length=200)

        # Simple keyword detection
        query_lower = query.lower()
        keywords_map = {
            "بطارية": "بطارية",
            "battery": "بطارية",
            "مولد": "مولد",
            "alternator": "مولد",
            "هواء": "MAF",
            "maf": "MAF",
            "كرنك": "الكرنك",
            "crank": "الكرنك",
            "أكسجين": "الأكسجين",
            "o2": "الأكسجين",
            "حرارة": "الحرارة",
            "temp": "الحرارة",
        }

        # Find matches quickly
        matches = []
        for comp in all_components:
            comp_ar = comp.get("componentAr", "").lower()
            comp_en = comp.get("componentEn", "").lower()

            for keyword, target in keywords_map.items():
                if keyword in query_lower and target.lower() in (
                    comp_ar + " " + comp_en
                ):
                    comp.pop("_id", None)
                    matches.append(comp)
                    break

        # If direct match found, return fast answer without AI
        if matches:
            first_match = matches[0]
            quick_answer = f"""## {first_match['componentAr']}

**الجهد الطبيعي:** {first_match['voltageNormal']} {first_match['unit']}

**النطاق المقبول:**
- الأدنى: {first_match['voltageMin']} {first_match['unit']}
- الأعلى: {first_match['voltageMax']} {first_match['unit']}

**طريقة القياس:** {first_match['measurementMethod']}

**الملاحظات:** {first_match['notes']}"""

            return {
                "answer": quick_answer,
                "matches": matches,
                "count": len(matches),
                "query": query,
                "fast_match": True,
            }

        # If no direct match, use AI (slower but more flexible)
        components_list = []
        for comp in all_components[:10]:  # Only first 10 for speed
            components_list.append(
                {
                    "المكون": comp.get("componentAr"),
                    "الجهد": comp.get("voltageNormal"),
                    "النطاق": f"{comp.get('voltageMin')}-{comp.get('voltageMax')} {comp.get('unit')}",
                }
            )

        llm = LlmChat(
            api_key=os.getenv("EMERGENT_LLM_KEY"),
            session_id=str(uuid.uuid4()),
            system_message="Answer voltage questions concisely in Arabic.",
        ).with_model("anthropic", "claude-sonnet-4-20250514")

        search_prompt = f"""سؤال: {query}
مراجع: {components_list}
أجب بإيجاز (50 كلمة): المكون، الجهد، النطاق، القياس."""

        response = await llm.send_message(UserMessage(text=search_prompt))
        response_text = response if isinstance(response, str) else response.text

        # Also find matching components
        matches = []
        query_lower = query.lower()
        keywords = [
            "هواء",
            "maf",
            "بطارية",
            "battery",
            "مولد",
            "alternator",
            "حساس",
            "sensor",
        ]

        for comp in all_components:
            comp_ar = comp.get("componentAr", "").lower()
            comp_en = comp.get("componentEn", "").lower()

            if query_lower in comp_ar or query_lower in comp_en:
                comp.pop("_id", None)
                matches.append(comp)
            elif any(
                kw in query_lower and kw in (comp_ar + " " + comp_en) for kw in keywords
            ):
                comp.pop("_id", None)
                matches.append(comp)

        return {
            "answer": response_text,
            "matches": matches,
            "count": len(matches),
            "query": query,
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
